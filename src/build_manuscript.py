"""Render the authoritative common manuscript JSON to Word and LaTeX.

No scientific text is inferred during conversion. PDF compilation is a separate
documented step; it requires Tectonic or a compatible LaTeX installation.
"""
from pathlib import Path
import json
import re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1]

def escape(text):
    mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}','β':r'\(\beta\)','θ':r'\(\theta\)','–':'--','—':'---'}
    return ''.join(mapping.get(c,c) for c in text)

def rows_for(source):
    if source=='candidate_scores':
        d=pd.read_csv(ROOT/'results/candidate_scores.tsv',sep='\t')
        return ['Gene','Signed score','Descending rank','Adjusted q'], [[r.gene,f'{r.published_score:.3f}',f'{int(r.protective_rank):,}',f'{r.reported_q:.3f}'] for _,r in d.iterrows()]
    d=pd.read_csv(ROOT/'results/raw_candidate_summary.csv').merge(pd.read_csv(ROOT/'results/raw_exploratory_NT_centered.csv'),on='gene',suffixes=('','_centered'))
    return ['Gene','Total-read contrast','NT-centered contrast'],[[r.gene,f'{r.descriptive_log2_CPM_contrast:.3f}',f'{r.NT_centered_descriptive_contrast:.3f}'] for _,r in d.iterrows()]

def set_borders(table):
    tblpr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    tblpr.append(borders)

def main(write_docx=True):
    data=json.loads((ROOT/'manuscript/paper_content.json').read_text(encoding='utf-8'))
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    styles=doc.styles
    for style in styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        styles[name].font.name='Calibri';styles[name].font.color.rgb=RGBColor(0,0,0)
    styles['Normal'].font.size=Pt(11)
    styles['Normal'].paragraph_format.space_after=Pt(6)
    styles['Normal'].paragraph_format.line_spacing=1.03
    styles['Title'].font.size=Pt(20)
    styles['Heading 1'].font.size=Pt(14)
    styles['Heading 1'].paragraph_format.space_before=Pt(9)
    styles['Heading 1'].paragraph_format.space_after=Pt(6)
    styles['Caption'].font.size=Pt(9)
    doc.core_properties.author=data['author'];doc.core_properties.last_modified_by=data['author']
    doc.core_properties.title=data['title'];doc.core_properties.subject='Fixed-set beta-cell immune-screen transfer and secretion evidence'
    doc.core_properties.comments=''
    doc.add_paragraph(data['title'],'Title')
    doc.add_paragraph(data['author']+' | '+data['date'])
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    tex=[r'\documentclass[11pt,letterpaper]{article}',r'\usepackage[margin=0.75in]{geometry}',
         r'\usepackage{graphicx,amsmath,amssymb,booktabs,float,array,tabularx,placeins}',r'\usepackage[unicode,hidelinks]{hyperref}',
         r'\usepackage[T1]{fontenc}',r'\usepackage{lmodern}',r'\setlength{\parindent}{0pt}',r'\setlength{\parskip}{6pt}',r'\setlength{\emergencystretch}{3em}',
         r'\renewcommand{\topfraction}{0.92}',r'\renewcommand{\bottomfraction}{0.85}',r'\renewcommand{\textfraction}{0.08}',r'\renewcommand{\floatpagefraction}{0.8}',
         r'\hypersetup{pdftitle={'+escape(data['title']).replace(r'\(\beta\)','β').replace(r'\(\theta\)','θ')+'},pdfauthor={'+escape(data['author'])+'}}',
         r'\begin{document}',r'{\LARGE\bfseries '+escape(data['title'])+r'\par}',
         r'\vspace{8pt}'+escape(data['author']+' | '+data['date'])+r'\par']
    for block in data['blocks']:
        typ=block['type']
        if typ=='page_break':
            doc.add_page_break();tex.append(r'\clearpage');continue
        if typ=='heading':
            level=block.get('level',1)
            if block['text']=='Discussion': tex.append(r'\FloatBarrier')
            doc.add_heading(block['text'],level)
            tex.append(('\\section*{' if level==1 else '\\subsection*{')+escape(block['text'])+'}');continue
        if typ in ['paragraph','reference']:
            p=doc.add_paragraph(block['text'])
            if typ=='reference':
                for r in p.runs:r.font.size=Pt(9)
                tex.append(r'{\small '+escape(block['text'])+r'\par}')
            else:tex.append(escape(block['text'])+r'\par')
            continue
        if typ=='equation':
            expr=block['latex'].replace(r'\tfrac12',r'\frac{1}{2}')
            fig=plt.figure(figsize=(7,.65));fig.text(.5,.5,'$'+expr+'$',ha='center',va='center',fontsize=15)
            image=ROOT/'manuscript/figures/estimand.png';fig.savefig(image,dpi=250,transparent=False,bbox_inches='tight',pad_inches=.08);plt.close(fig)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(image),width=Inches(6))
            tex.append(r'\['+expr+r'\]');continue
        if typ=='figure':
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(ROOT/'manuscript'/block['path']),width=Inches(6.85))
            doc.add_paragraph(block['caption'],'Caption')
            tex += [r'\begin{figure}[!htbp]\centering',r'\includegraphics[width=\linewidth]{'+block['path']+'}',r'\par\small '+escape(block['caption']),r'\end{figure}']
            continue
        if typ=='table':
            headers,rows=(block['headers'],block['rows']) if 'headers' in block else rows_for(block['source'])
            table=doc.add_table(rows=1,cols=len(headers));set_borders(table)
            table.autofit=False
            for cell,h in zip(table.rows[0].cells,headers):
                cell.text=h
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EAF0');cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs:run.bold=True
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                for cell,val in zip(table.add_row().cells,row):cell.text=val
            for row in table.rows:
                no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                        for run in p.runs:run.font.size=Pt(10)
            doc.add_paragraph(block['caption'],'Caption')
            wrapped='headers' in block
            tab_begin=(r'\begin{tabularx}{\linewidth}{@{}'+'>{\\raggedright\\arraybackslash}X'*len(headers)+r'@{}}') if wrapped else (r'\begin{tabular}{l'+'r'*(len(headers)-1)+'}')
            tex += [r'\begin{table}[!htbp]\centering\small',r'\renewcommand{\arraystretch}{1.15}',tab_begin,r'\toprule',' & '.join(map(escape,headers))+r'\\\midrule']
            tex += [' & '.join(map(escape,row))+r'\\' for row in rows]
            tex += [r'\bottomrule'+(r'\end{tabularx}' if wrapped else r'\end{tabular}'),r'\par\vspace{5pt}'+escape(block['caption']),r'\end{table}']
    if write_docx: doc.save(ROOT/'manuscript/paper.docx')
    tex.append(r'\end{document}')
    (ROOT/'manuscript/paper.tex').write_text('\n'.join(tex)+'\n',encoding='utf-8')
    print('Built Word and LaTeX from manuscript/paper_content.json')

if __name__=='__main__':main()
