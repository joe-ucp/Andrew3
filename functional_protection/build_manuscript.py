"""Render the common scientific text to editable Word.

The final PDF is exported from this same document using LibreOffice. This
renderer performs no inference and does not modify the foundation study.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent

def field(parent,name,value):
    x=OxmlElement(name)
    for k,v in value.items():x.set(qn(k),str(v))
    parent.append(x)
    return x

def main():
    data=json.loads((ROOT/'paper_content.json').read_text(encoding='utf-8'))
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.72);s.bottom_margin=Inches(.72)
    s.left_margin=Inches(1);s.right_margin=Inches(1)
    s.footer_distance=Inches(.32)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=d.styles[name];st.font.name='Cambria';st.font.color.rgb=RGBColor(0,0,0)
        for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    st=d.styles['Normal'];st.font.size=Pt(11);st.paragraph_format.line_spacing=1.06
    st.paragraph_format.space_after=Pt(6);st.paragraph_format.widow_control=True
    d.styles['Title'].font.size=Pt(20);d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(7)
    d.styles['Subtitle'].font.size=Pt(11);d.styles['Subtitle'].font.italic=True
    d.styles['Subtitle'].paragraph_format.space_after=Pt(8)
    for name,size in [('Heading 1',14),('Heading 2',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.italic=False
    d.styles['Caption'].paragraph_format.space_after=Pt(8)
    props=d.core_properties;props.title=data['title'];props.author=data['author']
    props.last_modified_by=data['author'];props.subject=data['subtitle'];props.comments=''
    d.add_paragraph(data['title'],'Title')
    d.add_paragraph(data['subtitle'],'Subtitle')
    p=d.add_paragraph(data['author']+' | '+data['date']);p.paragraph_format.space_after=Pt(10)
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(p._p,'w:fldSimple',{'w:instr':'PAGE'})
    for b in data['blocks']:
        kind=b['type']
        if kind=='heading':
            d.add_heading(b['text'],b.get('level',1))
        elif kind in ('paragraph','reference','equation'):
            p=d.add_paragraph(b['text'])
            if kind=='reference':
                p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(9.5)
            if kind=='equation':
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
                for r in p.runs:r.font.size=Pt(10.5)
        elif kind=='figure':
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(ROOT/b['path']),width=Inches(6.5))
            d.add_paragraph(b['caption'],'Caption')
        elif kind=='table':
            # Caption above the table remains attached to the first row.
            p=d.add_paragraph(b['caption'],'Caption');p.paragraph_format.keep_with_next=True
            t=d.add_table(rows=1,cols=len(b['headers']));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
            for col,width in zip(t.columns,b['widths']):col.width=Inches(width)
            borders=field(t._tbl.tblPr,'w:tblBorders',{})
            for edge in ['top','left','bottom','right','insideH','insideV']:
                field(borders,'w:'+edge,{'w:val':'single','w:sz':'4','w:color':'D9D9D9'})
            field(t.rows[0]._tr.get_or_add_trPr(),'w:tblHeader',{})
            for cell,text,width in zip(t.rows[0].cells,b['headers'],b['widths']):
                cell.text=text;cell.width=Inches(width)
                field(cell._tc.get_or_add_tcPr(),'w:shd',{'w:fill':'E9EEF0'})
                for r in cell.paragraphs[0].runs:r.bold=True
            for values in b['rows']:
                row=t.add_row()
                for cell,text,width in zip(row.cells,values,b['widths']):cell.text=text;cell.width=Inches(width)
            for row in t.rows:
                field(row._tr.get_or_add_trPr(),'w:cantSplit',{})
                for cell in row.cells:
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    m=field(cell._tc.get_or_add_tcPr(),'w:tcMar',{})
                    for edge in ['top','left','bottom','right']:field(m,'w:'+edge,{'w:w':'85','w:type':'dxa'})
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                        p.paragraph_format.line_spacing=1.0
                        for r in p.runs:r.font.size=Pt(9.5)
            d.add_paragraph().paragraph_format.space_after=Pt(1)
    d.save(ROOT/'paper.docx')
    print('Created functional_protection/paper.docx')

if __name__=='__main__':main()
